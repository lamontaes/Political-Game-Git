"""Compare the Census item-code sums behind service-charges-2022.json with
Table 1's own sub-lines, one service at a time, so a gap against Table 1 is
traced to the line that carries it (pool item gov-service-charges-table1-gaps).

Reads data/source/government-finances/census-tables/22statetypepu.txt (item
codes, thousands of dollars) and 22slsstab1.xlsx (Table 1, thousands). Table 1's
footnote 1 says duplicative intergovernmental transactions are excluded; the
item-code file carries them gross.

Run: python3 scripts/research/compare-service-charges-table1.py
"""

import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
CENSUS = ROOT / "data/source/government-finances/census-tables"

# The item codes (technical documentation, "Item Codes and Short Descriptions")
# against the Table 1 sub-line they should fill.
LINES = {
    "Education": (41, ["A09", "A10", "A12", "A16", "A18", "A21"]),
    "Hospitals": (44, ["A36"]),
    "Highways": (45, ["A44", "A45"]),
    "Airports": (46, ["A01"]),
    "Parking": (47, ["A60"]),
    "Ports": (48, ["A87"]),
    "Natural resources": (49, ["A59"]),
    "Parks": (50, ["A61"]),
    "Housing": (51, ["A50"]),
    "Sewerage": (52, ["A80"]),
    "Solid waste": (53, ["A81"]),
    "Other charges": (54, ["A03", "A89"]),
    "Interest earnings": (57, ["U20"]),
    "Special assessments": (58, ["U01"]),
    "Sale of property": (59, ["U11"]),
    "Other general revenue": (60, ["U30", "U40", "U41", "U50", "U95", "U99"]),
}
# The states and levels whose current charges or miscellaneous revenue differ
# from Table 1 by more than 10 percent: (state FIPS, Table 1 name, level code).
ROWS = [
    ("15", "Hawaii", "3"), ("21", "Kentucky", "3"), ("29", "Missouri", "3"),
    ("40", "Oklahoma", "2"), ("45", "South Carolina", "3"),
    ("47", "Tennessee", "2"), ("47", "Tennessee", "3"), ("48", "Texas", "2"),
]


def main() -> None:
    amounts: dict[tuple[str, str], dict[str, int]] = {}
    for line in (CENSUS / "22statetypepu.txt").read_text().splitlines():
        if len(line) >= 20:
            amounts.setdefault((line[0:2], line[2]), {})[line[4:7]] = int(line[8:20])
    wb = openpyxl.load_workbook(CENSUS / "22slsstab1.xlsx", read_only=True, data_only=True)
    sheet = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    first = {c.strip(): j for j, c in enumerate(sheet[8]) if c and j > 1}
    for fips, name, level in ROWS:
        items = amounts[(fips, level)]
        column = first[name] + (2 if level == "2" else 3)
        print(f"{name} {'state' if level == '2' else 'local'} (thousands of dollars)")
        for line, (row, codes) in LINES.items():
            coded = sum(items.get(c, 0) for c in codes)
            table = sheet[row][column] or 0
            if abs(coded - table) > max(5000, 0.1 * abs(table)):
                print(f"  {line:<22} codes {coded:>10,}  Table 1 {table:>10,}  Table 1 minus codes {table - coded:>+10,}  ({'+'.join(codes)})")


if __name__ == "__main__":
    main()
