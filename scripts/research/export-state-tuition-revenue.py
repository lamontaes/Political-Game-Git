"""Export what a state tuition freeze holds back, for each state, to JSON.

Reads the Census Bureau's 2022 state and local finance table
(data/source/government-finances/census-tables/22slsstab1.xlsx) and SHEEO's
SHEF report data (data/source/higher-education-finance/, README there) and
writes data/research/money/state-tuition-revenue.json:

- for each state, the share of the state government's current charges (the
  budget's charges and fees, Census line 23) that is tuition: SHEF's gross
  tuition and fee revenue at public colleges in fiscal 2022, at most the
  state's own charges at institutions of higher education (Census line 25,
  which also holds dorm and dining charges), over line 23;
- how fast tuition rises in a year no freeze holds it: the growth of gross
  tuition and fee revenue per full-time student, fiscal 2015 to 2025, for the
  nation (the central value) and each state (the spread).

D.C. is left out: its own budget keeps no state-level books. The territories
are in neither source.

Run: python3 scripts/research/export-state-tuition-revenue.py
Needs openpyxl.
"""

import hashlib
import json
import pathlib
import statistics

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[2]
CENSUS = ROOT / "data/source/government-finances/census-tables/22slsstab1.xlsx"
CENSUS_SHA256 = "4dd123c5ad520e33d9692f622d389a020e2d3a7e6e782763663854dd84643b9b"
SHEF = ROOT / "data/source/higher-education-finance/SHEEO_SHEF_FY25_Report_Data.xlsx"
SHEF_SHA256 = "3a48fd92b392999800016b378b4016b9dfff13c5d268d0004ec4b219f8d74860"
OUT = ROOT / "data/research/money/state-tuition-revenue.json"

CURRENT_CHARGES_LINE = 23
HIGHER_EDUCATION_CHARGES_LINE = 25
SHARE_YEAR = 2022
GROWTH_FROM, GROWTH_TO = 2015, 2025

STATES = {
    "Alabama": "US-AL", "Alaska": "US-AK", "Arizona": "US-AZ",
    "Arkansas": "US-AR", "California": "US-CA", "Colorado": "US-CO",
    "Connecticut": "US-CT", "Delaware": "US-DE", "Florida": "US-FL",
    "Georgia": "US-GA", "Hawaii": "US-HI", "Idaho": "US-ID",
    "Illinois": "US-IL", "Indiana": "US-IN", "Iowa": "US-IA",
    "Kansas": "US-KS", "Kentucky": "US-KY", "Louisiana": "US-LA",
    "Maine": "US-ME", "Maryland": "US-MD", "Massachusetts": "US-MA",
    "Michigan": "US-MI", "Minnesota": "US-MN", "Mississippi": "US-MS",
    "Missouri": "US-MO", "Montana": "US-MT", "Nebraska": "US-NE",
    "Nevada": "US-NV", "New Hampshire": "US-NH", "New Jersey": "US-NJ",
    "New Mexico": "US-NM", "New York": "US-NY", "North Carolina": "US-NC",
    "North Dakota": "US-ND", "Ohio": "US-OH", "Oklahoma": "US-OK",
    "Oregon": "US-OR", "Pennsylvania": "US-PA", "Rhode Island": "US-RI",
    "South Carolina": "US-SC", "South Dakota": "US-SD", "Tennessee": "US-TN",
    "Texas": "US-TX", "Utah": "US-UT", "Vermont": "US-VT",
    "Virginia": "US-VA", "Washington": "US-WA", "West Virginia": "US-WV",
    "Wisconsin": "US-WI", "Wyoming": "US-WY",
}


def checked(path, expected):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != expected:
        raise SystemExit(f"Source hash mismatch for {path}: {digest}")
    return digest


def census_state_charges():
    """Census line 23 and 25 for each state government, in dollars."""
    sheet = openpyxl.load_workbook(CENSUS, read_only=True, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[8]
    by_line = {row[0]: row for row in rows if isinstance(row[0], int)}
    out = {}
    for index, name in enumerate(header):
        if isinstance(name, str) and name.strip() in STATES:
            state_column = index + 2
            out[STATES[name.strip()]] = {
                "currentCharges": int(by_line[CURRENT_CHARGES_LINE][state_column]) * 1000,
                "higherEducationCharges": int(by_line[HIGHER_EDUCATION_CHARGES_LINE][state_column]) * 1000,
            }
    if len(out) != 50:
        raise SystemExit(f"Expected 50 states in the Census table, found {len(out)}")
    return out


def shef_rows():
    sheet = openpyxl.load_workbook(SHEF, read_only=True, data_only=True)["Report Data"]
    rows = list(sheet.iter_rows(values_only=True))
    column = {name: index for index, name in enumerate(rows[0])}
    by_state = {}
    for row in rows[1:]:
        by_state.setdefault(row[column["State"]], {})[row[column["FY"]]] = {
            "grossTuition": row[column["Gross Tuition and Fee Revenue"]],
            "fte": row[column["Net FTE Enrollment"]],
        }
    return by_state


def yearly_growth(years):
    before = years[GROWTH_FROM]["grossTuition"] / years[GROWTH_FROM]["fte"]
    after = years[GROWTH_TO]["grossTuition"] / years[GROWTH_TO]["fte"]
    return (after / before) ** (1 / (GROWTH_TO - GROWTH_FROM)) - 1


def main():
    census_digest = checked(CENSUS, CENSUS_SHA256)
    shef_digest = checked(SHEF, SHEF_SHA256)
    charges = census_state_charges()
    shef = shef_rows()
    places = {}
    growths = []
    for name, code in sorted(STATES.items(), key=lambda item: item[1]):
        years = shef[name]
        gross = int(round(years[SHARE_YEAR]["grossTuition"]))
        census = charges[code]
        tuition = min(gross, census["higherEducationCharges"])
        growth = yearly_growth(years)
        growths.append(growth)
        places[code] = {
            "name": name,
            "currentCharges2022": census["currentCharges"],
            "higherEducationCharges2022": census["higherEducationCharges"],
            "grossTuitionAndFees2022": gross,
            "tuitionShareOfCharges": round(tuition / census["currentCharges"], 4),
            "tuitionGrowthPerYear2015to2025": round(growth, 4),
        }
    quartiles = statistics.quantiles(growths, n=4)
    out = {
        "id": "state-tuition-revenue",
        "asOf": "2025",
        "sourceKind": "official",
        "sources": [
            {
                "title": "State and Local Government Finances by Level of Government and by State: 2022 (Table 1), lines 23 and 25, state government",
                "publisher": "U.S. Census Bureau, 2022 Census of Governments: Finance",
                "file": str(CENSUS.relative_to(ROOT)),
                "sha256": census_digest,
            },
            {
                "title": "State Higher Education Finance: FY 2025, report data (unadjusted)",
                "publisher": "State Higher Education Executive Officers Association",
                "url": "https://shef.sheeo.org/wp-content/uploads/2026/04/SHEEO_SHEF_FY25_Report_Data.xlsx",
                "file": str(SHEF.relative_to(ROOT)),
                "sha256": shef_digest,
            },
        ],
        "script": "scripts/research/export-state-tuition-revenue.py",
        "method": "tuitionShareOfCharges: SHEF gross tuition and fee revenue at public colleges in fiscal 2022, at most the state government's charges at institutions of higher education (Census line 25, which also holds dorm and dining charges), over the state government's current charges (Census line 23). Tuition growth: gross tuition and fee revenue per net full-time student, fiscal 2015 to 2025, as a yearly rate, nominal.",
        "tuitionGrowthPerYear": {
            "central": round(yearly_growth(shef["U.S."]), 4),
            "low": round(quartiles[0], 4),
            "high": round(quartiles[2], 4),
            "basis": "ESTIMATED FROM AVERAGE: the nation's growth is the central value; the spread is the states' interquartile range.",
        },
        "places": places,
    }
    OUT.write_text(json.dumps(out, indent=2) + "\n")
    print(f"Wrote {OUT.relative_to(ROOT)}: {len(places)} states")


if __name__ == "__main__":
    main()
